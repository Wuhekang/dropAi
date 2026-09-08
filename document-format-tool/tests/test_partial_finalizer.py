from __future__ import annotations

from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, RGBColor
from lxml import etree

from word_formatter.core import finalizer
from word_formatter.core.finalizer import finalize_docx


class PartialFinalizerTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix="partial-finalizer-test-")
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / "reviewed.docx"
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(8)
        run = paragraph.add_run("必须保留的作者正文")
        run.font.color.rgb = RGBColor(255, 0, 0)
        document.add_comment(run, text="审阅备注", author="Reviewer")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), " REF missing ")
        field_run = OxmlElement("w:r")
        field_text = OxmlElement("w:t")
        field_text.text = "6"
        field_run.append(field_text)
        field.append(field_run)
        paragraph._p.append(field)
        header = document.sections[0].header.paragraphs[0]
        header.paragraph_format.left_indent = Cm(8)
        header.add_run("页眉正文").font.color.rgb = RGBColor(255, 0, 0)
        document.save(self.path)
        self.original = self.path.read_bytes()

    def xml(self, part="word/document.xml"):
        with zipfile.ZipFile(self.path) as package:
            return etree.fromstring(package.read(part))

    def comment_parts(self):
        with zipfile.ZipFile(self.path) as package:
            return {name: package.read(name) for name in package.namelist() if "comment" in name.lower()}

    def assert_safe_error(self, errors, label, kind):
        self.assertEqual(errors, [{
            "item": f"最终清理：{label}",
            "reason": f"{kind}：该项处理失败，已保留原有内容",
        }])
        self.assertNotIn("private", str(errors))

    def test_namespace_failure_rolls_back_reference_action_and_runs_all_later_cleanup(self):
        original_lock = finalizer._lock_unresolved_references

        def fail_document(root, bookmarks):
            if root.tag == qn("w:document"):
                next(root.iter(qn("w:t"))).text = "CORRUPTED private document text"
                root.xpath(".//unbound:field")
            return original_lock(root, bookmarks)

        errors = []
        with patch.object(finalizer, "_lock_unresolved_references", side_effect=fail_document):
            stats = finalize_docx(self.path, errors=errors)

        self.assert_safe_error(errors, "引用保护", "XPathEvalError")
        root = self.xml()
        self.assertEqual(next(root.iter(qn("w:t"))).text, "必须保留的作者正文")
        self.assertIsNone(next(root.iter(qn("w:fldSimple"))).get(qn("w:fldLock")))
        self.assertEqual(next(root.iter(qn("w:ind"))).get(qn("w:left")), "0")
        self.assertEqual(next(root.iter(qn("w:color"))).get(qn("w:val")), "000000")
        self.assertFalse(list(root.iter(qn("w:commentReference"))))
        self.assertFalse(self.comment_parts())
        self.assertEqual(stats["unresolved_references_locked"], 0)
        self.assertGreater(stats["unsafe_indents_reset"], 0)
        self.assertGreater(stats["comment_markup_removed"], 0)
        self.assertGreater(stats["comment_parts_removed"], 0)
        self.assertGreater(stats["red_fonts_blackened"], 0)
        self.assertTrue(all(type(value) is int for value in stats.values()))
        self.assertEqual(Document(self.path).paragraphs[0].text, "必须保留的作者正文")

    def test_failed_indent_action_keeps_prior_reference_and_later_comment_color_actions(self):
        original_apply = finalizer.IndentGuard.apply

        def fail_document(guard, root):
            if root.tag == qn("w:document"):
                next(root.iter(qn("w:t"))).text = "CORRUPTED private text"
                raise ValueError("private token or path")
            return original_apply(guard, root)

        errors = []
        with patch.object(finalizer.IndentGuard, "apply", new=fail_document):
            stats = finalize_docx(self.path, errors)
        self.assert_safe_error(errors, "缩进保护", "ValueError")
        root = self.xml()
        self.assertEqual(next(root.iter(qn("w:t"))).text, "必须保留的作者正文")
        self.assertEqual(next(root.iter(qn("w:fldSimple"))).get(qn("w:fldLock")), "true")
        self.assertEqual(next(root.iter(qn("w:ind"))).get(qn("w:left")), str(Cm(8).twips))
        self.assertEqual(next(root.iter(qn("w:color"))).get(qn("w:val")), "000000")
        self.assertFalse(self.comment_parts())
        header = self.xml("word/header1.xml")
        self.assertEqual(next(header.iter(qn("w:ind"))).get(qn("w:left")), "0")
        self.assertGreater(stats["unsafe_indents_reset"], 0)

    def test_failed_comment_markup_keeps_comment_package_coherent_and_continues_color(self):
        original_parts = self.comment_parts()
        with zipfile.ZipFile(self.path) as package:
            original_rels = package.read("word/_rels/document.xml.rels")
            original_types = package.read("[Content_Types].xml")
        original_remove = finalizer._remove_comment_markup

        def fail_document(root):
            if root.tag == qn("w:document"):
                original_remove(root)
                raise RuntimeError("private review text")
            return original_remove(root)

        errors = []
        with patch.object(finalizer, "_remove_comment_markup", side_effect=fail_document):
            stats = finalize_docx(self.path, errors)
        self.assert_safe_error(errors, "批注清理", "RuntimeError")
        root = self.xml()
        self.assertTrue(list(root.iter(qn("w:commentReference"))))
        self.assertEqual(next(root.iter(qn("w:color"))).get(qn("w:val")), "000000")
        self.assertEqual(next(root.iter(qn("w:ind"))).get(qn("w:left")), "0")
        self.assertEqual(self.comment_parts(), original_parts)
        with zipfile.ZipFile(self.path) as package:
            self.assertEqual(package.read("word/_rels/document.xml.rels"), original_rels)
            self.assertEqual(package.read("[Content_Types].xml"), original_types)
        self.assertEqual(stats["comment_parts_removed"], 0)
        self.assertEqual(stats["comment_markup_removed"], 0)
        Document(self.path)

    def test_failed_red_cleanup_rolls_back_only_that_part_and_later_header_still_runs(self):
        original_blacken = finalizer._blacken_red_fonts

        def fail_document(root):
            if root.tag == qn("w:document"):
                original_blacken(root)
                next(root.iter(qn("w:t"))).text = "CORRUPTED"
                raise RuntimeError("private text")
            return original_blacken(root)

        errors = []
        with patch.object(finalizer, "_blacken_red_fonts", side_effect=fail_document):
            stats = finalize_docx(self.path, errors)
        self.assert_safe_error(errors, "红色字体清理", "RuntimeError")
        root = self.xml()
        self.assertEqual(next(root.iter(qn("w:t"))).text, "必须保留的作者正文")
        self.assertEqual(next(root.iter(qn("w:color"))).get(qn("w:val")), "FF0000")
        self.assertEqual(next(root.iter(qn("w:ind"))).get(qn("w:left")), "0")
        self.assertFalse(self.comment_parts())
        header = self.xml("word/header1.xml")
        self.assertEqual(next(header.iter(qn("w:color"))).get(qn("w:val")), "000000")
        self.assertGreater(stats["red_fonts_blackened"], 0)

    def test_failed_comment_content_types_discards_all_metadata_removal(self):
        original_parts = self.comment_parts()
        with zipfile.ZipFile(self.path) as package:
            original_rels = package.read("word/_rels/document.xml.rels")
            original_types = package.read("[Content_Types].xml")

        def fail_after_mutation(root):
            root.clear()
            raise ValueError("private metadata")

        errors = []
        with patch.object(finalizer, "_remove_comment_content_types", side_effect=fail_after_mutation):
            stats = finalize_docx(self.path, errors)
        self.assert_safe_error(errors, "批注类型清理", "ValueError")
        self.assertEqual(self.comment_parts(), original_parts)
        with zipfile.ZipFile(self.path) as package:
            self.assertEqual(package.read("word/_rels/document.xml.rels"), original_rels)
            self.assertEqual(package.read("[Content_Types].xml"), original_types)
        self.assertEqual(stats["comment_parts_removed"], 0)
        self.assertGreater(stats["comment_markup_removed"], 0)
        self.assertFalse(list(self.xml().iter(qn("w:commentReference"))))
        Document(self.path)

    def test_no_error_collector_retains_fail_fast_and_original_file(self):
        with patch.object(finalizer, "_lock_unresolved_references", side_effect=ValueError("private error")):
            with self.assertRaises(ValueError):
                finalize_docx(self.path)
        self.assertEqual(self.path.read_bytes(), self.original)

    def test_fatal_errors_are_not_reported_as_partial_success(self):
        for exc in (OSError(28, "disk full"), zipfile.BadZipFile("corrupt package"), MemoryError(), KeyboardInterrupt()):
            with self.subTest(kind=type(exc).__name__):
                errors = []
                with patch.object(finalizer, "_lock_unresolved_references", side_effect=exc):
                    with self.assertRaises(type(exc)):
                        finalize_docx(self.path, errors)
                self.assertFalse(errors)
                self.assertEqual(self.path.read_bytes(), self.original)

    def test_invalid_required_xml_is_fatal(self):
        with zipfile.ZipFile(self.path) as package:
            original_parts = {name: package.read(name) for name in package.namelist()}
        for part in ("word/document.xml", "[Content_Types].xml", "_rels/.rels", "word/_rels/document.xml.rels"):
            with self.subTest(part=part):
                with zipfile.ZipFile(self.path, "w") as package:
                    for name, data in original_parts.items():
                        package.writestr(name, b"<broken" if name == part else data)
                corrupted = self.path.read_bytes()
                errors = []
                with self.assertRaises(etree.XMLSyntaxError):
                    finalize_docx(self.path, errors)
                self.assertFalse(errors)
                self.assertEqual(self.path.read_bytes(), corrupted)

    def test_disk_write_failure_preserves_input_and_cleans_temporary(self):
        with patch.object(zipfile.ZipFile, "writestr", side_effect=OSError(28, "disk full")):
            with self.assertRaises(OSError):
                finalize_docx(self.path, [])
        self.assertEqual(self.path.read_bytes(), self.original)
        self.assertFalse(list(self.path.parent.glob("*.finalizing.docx")))


if __name__ == "__main__":
    unittest.main()
