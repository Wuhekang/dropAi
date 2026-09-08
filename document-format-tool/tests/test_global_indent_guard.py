from __future__ import annotations

import base64
from dataclasses import fields
from io import BytesIO
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree


TOOL_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(TOOL_ROOT))

from word_formatter.core.finalizer import finalize_docx  # noqa: E402
from word_formatter.core.processor import DocumentProcessor  # noqa: E402
from word_formatter.models.rules import (  # noqa: E402
    DocumentRules,
    ParagraphRule,
    enforce_safe_indentation,
)


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}
INDENT_FIELDS = (
    "left_indent_cm", "right_indent_cm", "left_indent_chars",
    "right_indent_chars", "first_line_indent_chars", "special_indent_chars",
)
TWIP_ATTRIBUTES = ("left", "right", "start", "end", "firstLine", "hanging")
CHAR_ATTRIBUTES = tuple(f"{name}Chars" for name in TWIP_ATTRIBUTES)
PIXEL_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aP1cAAAAASUVORK5CYII="
)


def _ind(parent, **attributes):
    """Attach explicit XML without python-docx coercing malformed input."""
    node = parent.find(qn("w:ind"))
    if node is None:
        node = OxmlElement("w:ind")
        parent.append(node)
    for name, value in attributes.items():
        node.set(qn(f"w:{name}"), str(value))
    return node


def _paragraph_ind(paragraph, **attributes):
    return _ind(paragraph._p.get_or_add_pPr(), **attributes)


def _raw_paragraph(text, **attributes):
    paragraph = OxmlElement("w:p")
    properties = OxmlElement("w:pPr")
    paragraph.append(properties)
    _ind(properties, **attributes)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "24")
    run_properties.append(size)
    run.append(run_properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    paragraph.append(run)
    return paragraph


def _package(path):
    with ZipFile(path) as package:
        return {name: package.read(name) for name in package.namelist()}


def _replace_parts(path, overrides):
    parts = _package(path)
    parts.update(overrides)
    with ZipFile(path, "w", ZIP_DEFLATED) as package:
        for name, data in parts.items():
            package.writestr(name, data)


def _root(path, part="word/document.xml"):
    with ZipFile(path) as package:
        return etree.fromstring(package.read(part))


def _paragraph_by_text(root, text):
    return next(
        p for p in root.xpath(".//w:p", namespaces=NS)
        if "".join(p.xpath(".//w:t/text()", namespaces=NS)) == text
    )


def _semantic_snapshot(parts):
    """Compare author content and live objects, independently of formatting."""
    snapshot = {}
    for name, data in parts.items():
        if name.startswith("word/media/") or name.endswith(".rels"):
            snapshot[name] = data
        elif name.startswith("word/") and name.endswith(".xml"):
            root = etree.fromstring(data)
            snapshot[name] = (
                root.xpath(".//w:t/text()", namespaces=NS),
                root.xpath(".//w:instrText/text()", namespaces=NS),
                [dict(node.attrib) for node in root.xpath(
                    ".//w:fldChar | .//w:fldSimple | .//w:bookmarkStart | .//w:bookmarkEnd",
                    namespaces=NS,
                )],
                [etree.tostring(node) for node in root.xpath(".//w:drawing", namespaces=NS)],
            )
    return snapshot


class RuleIndentGuardTests(unittest.TestCase):
    def test_every_paragraph_rule_is_guarded_even_when_disabled(self):
        rules = DocumentRules()
        count = 0
        for item in fields(rules):
            rule = getattr(rules, item.name)
            if isinstance(rule, ParagraphRule):
                rule.enabled = False
                for name in INDENT_FIELDS:
                    setattr(rule, name, 999.0)
                    count += 1
        self.assertEqual(enforce_safe_indentation(rules), count)
        for item in fields(rules):
            rule = getattr(rules, item.name)
            if isinstance(rule, ParagraphRule):
                for name in INDENT_FIELDS:
                    self.assertEqual(getattr(rule, name), 0.0, f"{item.name}.{name}")
        self.assertEqual(enforce_safe_indentation(rules), 0)

    def test_negative_nonfinite_invalid_and_out_of_range_values_reset(self):
        for name in INDENT_FIELDS:
            limit = 2.0 if name.endswith("_cm") else 4.0
            for value in (-0.01, float("nan"), float("inf"), -float("inf"), limit + 0.001, "999", None, True):
                with self.subTest(field=name, value=value):
                    rules = DocumentRules()
                    setattr(rules.heading_4, name, value)
                    self.assertEqual(enforce_safe_indentation(rules), 1)
                    self.assertEqual(getattr(rules.heading_4, name), 0.0)
                    self.assertEqual(enforce_safe_indentation(rules), 0)

    def test_safe_body_toc_reference_and_exact_cm_boundaries_survive(self):
        rules = DocumentRules()
        rules.normal_text.left_indent_cm = 2.0
        rules.normal_text.right_indent_cm = 2.0
        rules.toc_3.font_size_pt = 12.0
        rules.toc_3.left_indent_chars = 4.0
        rules.reference.special_indent_mode = "hanging"
        rules.reference.special_indent_chars = 2.0
        before = rules.to_dict()
        self.assertEqual(enforce_safe_indentation(rules), 0)
        self.assertEqual(rules.to_dict(), before)

    def test_character_limits_also_use_the_rules_font_size(self):
        for name in (name for name in INDENT_FIELDS if name.endswith("_chars")):
            with self.subTest(field=name):
                rules = DocumentRules()
                rules.heading_4.font_size_pt = 36.0
                setattr(rules.heading_4, name, 2.0)  # 2 characters = 2.54 cm.
                self.assertEqual(enforce_safe_indentation(rules), 1)
                self.assertEqual(getattr(rules.heading_4, name), 0.0)


class PackageIndentGuardTests(unittest.TestCase):
    def test_every_indent_alias_rejects_negative_nonfinite_and_oversize_xml(self):
        with tempfile.TemporaryDirectory(prefix="indent_aliases_") as directory:
            path = Path(directory) / "aliases.docx"
            document = Document()
            expected = []
            for name in TWIP_ATTRIBUTES + CHAR_ATTRIBUTES:
                limit = 400 if name.endswith("Chars") else Cm(2).twips
                for index, value in enumerate(("-1", "NaN", "INF", "-INF", "bad", str(limit + 1), "999999")):
                    text = f"{name}-{index}"
                    paragraph = document.add_paragraph(text)
                    paragraph.runs[0].font.size = Pt(12)
                    _paragraph_ind(paragraph, **{name: value})
                    expected.append((text, name))
            document.save(path)
            stats = finalize_docx(path)
            self.assertGreaterEqual(stats["unsafe_indents_reset"], len(expected))
            root = _root(path)
            for text, name in expected:
                paragraph = _paragraph_by_text(root, text)
                node = paragraph.find("w:pPr/w:ind", NS)
                self.assertEqual(node.get(qn(f"w:{name}")), "0", (text, name))
            self.assertEqual(finalize_docx(path)["unsafe_indents_reset"], 0)

    def test_all_story_parts_nested_content_and_table_offsets_are_guarded(self):
        with tempfile.TemporaryDirectory(prefix="indent_stories_") as directory:
            path = Path(directory) / "stories.docx"
            document = Document()
            _paragraph_ind(document.add_paragraph("封面作者信息"), left=9000)
            _paragraph_ind(document.add_paragraph("独创性声明原文"), endChars=900)
            sdt = OxmlElement("w:sdt")
            content = OxmlElement("w:sdtContent")
            content.append(_raw_paragraph("内容控件目录", start=9000))
            sdt.append(content)
            document.element.body.insert(0, sdt)
            table = document.add_table(1, 1)
            table.cell(0, 0).text = "表格单元格原文"
            _paragraph_ind(table.cell(0, 0).paragraphs[0], right=9000)
            table_indent = OxmlElement("w:tblInd")
            table_indent.set(qn("w:type"), "dxa")
            table_indent.set(qn("w:w"), "9000")
            table._tbl.tblPr.append(table_indent)
            textbox = OxmlElement("w:txbxContent")
            textbox.append(_raw_paragraph("文本框原文", leftChars=900))
            pict = OxmlElement("w:pict")
            shape = etree.SubElement(pict, "{urn:schemas-microsoft-com:vml}shape")
            text_box = etree.SubElement(shape, "{urn:schemas-microsoft-com:vml}textbox")
            text_box.append(textbox)
            document.add_paragraph().add_run()._r.append(pict)
            header = document.sections[0].header.paragraphs[0]
            header.text = "页眉原文"
            _paragraph_ind(header, rightChars=900)
            footer = document.sections[0].footer.paragraphs[0]
            footer.text = "页脚原文"
            _paragraph_ind(footer, hanging=9000)
            document.save(path)
            extra = {}
            for part, tag, text in (
                ("word/footnotes.xml", "footnotes", "脚注原文"),
                ("word/endnotes.xml", "endnotes", "尾注原文"),
                ("word/glossary/document.xml", "document", "构建基块原文"),
            ):
                root = etree.Element(qn(f"w:{tag}"), nsmap={"w": W_NS})
                root.append(_raw_paragraph(text, firstLineChars=900))
                extra[part] = etree.tostring(root)
            _replace_parts(path, extra)
            before = _semantic_snapshot(_package(path))
            self.assertGreaterEqual(finalize_docx(path)["unsafe_indents_reset"], 11)
            parts = _package(path)
            self.assertEqual(_semantic_snapshot(parts), before)
            for part, data in parts.items():
                if not (part.startswith("word/") and part.endswith(".xml")):
                    continue
                root = etree.fromstring(data)
                for node in root.xpath(".//w:ind", namespaces=NS):
                    for name in TWIP_ATTRIBUTES + CHAR_ATTRIBUTES:
                        value = node.get(qn(f"w:{name}"))
                        if value is not None:
                            self.assertLessEqual(float(value), 400 if name.endswith("Chars") else Cm(2).twips, (part, name, value))
                            self.assertGreaterEqual(float(value), 0, (part, name, value))
            self.assertEqual(_root(path).find(".//w:tblInd", NS).get(qn("w:w")), "0")

    def test_style_inheritance_and_numbering_definitions_cannot_hide_bad_indents(self):
        with tempfile.TemporaryDirectory(prefix="indent_styles_") as directory:
            path = Path(directory) / "styles.docx"
            document = Document()
            base = document.styles.add_style("UnsafeBase", WD_STYLE_TYPE.PARAGRAPH)
            base.font.size = Pt(36)
            _ind(base._element.get_or_add_pPr(), left=9000, startChars=200)
            child = document.styles.add_style("InheritedChild", WD_STYLE_TYPE.PARAGRAPH)
            child.base_style = base
            document.add_paragraph("继承样式原文", child)
            numbering = document.part.numbering_part.element
            abstract = OxmlElement("w:abstractNum")
            abstract.set(qn("w:abstractNumId"), "900")
            level = OxmlElement("w:lvl")
            level.set(qn("w:ilvl"), "0")
            properties = OxmlElement("w:pPr")
            _ind(properties, end=9000, hangingChars=500)
            level.append(properties)
            abstract.append(level)
            numbering.append(abstract)
            document.save(path)
            self.assertGreaterEqual(finalize_docx(path)["unsafe_indents_reset"], 4)
            styles = _root(path, "word/styles.xml")
            safe_base = styles.xpath(".//w:style[@w:styleId='UnsafeBase']/w:pPr/w:ind", namespaces=NS)[0]
            self.assertEqual(safe_base.get(qn("w:left")), "0")
            self.assertEqual(safe_base.get(qn("w:startChars")), "0")
            self.assertEqual(styles.xpath(".//w:style[@w:styleId='InheritedChild']/w:basedOn/@w:val", namespaces=NS), ["UnsafeBase"])
            safe_level = _root(path, "word/numbering.xml").xpath(".//w:abstractNum[@w:abstractNumId='900']//w:ind", namespaces=NS)[0]
            self.assertEqual(safe_level.get(qn("w:end")), "0")
            self.assertEqual(safe_level.get(qn("w:hangingChars")), "0")

    def test_inherited_style_font_size_limits_direct_character_indents(self):
        with tempfile.TemporaryDirectory(prefix="indent_font_") as directory:
            path = Path(directory) / "font.docx"
            document = Document()
            base = document.styles.add_style("LargeFontBase", WD_STYLE_TYPE.PARAGRAPH)
            base.font.size = Pt(36)
            child = document.styles.add_style("LargeFontChild", WD_STYLE_TYPE.PARAGRAPH)
            child.base_style = base
            paragraph = document.add_paragraph("继承三十六磅字号", child)
            _paragraph_ind(paragraph, leftChars=200, firstLineChars=200)
            document.save(path)
            finalize_docx(path)
            node = _paragraph_by_text(_root(path), "继承三十六磅字号").find("w:pPr/w:ind", NS)
            self.assertEqual(node.get(qn("w:leftChars")), "0")
            self.assertEqual(node.get(qn("w:firstLineChars")), "0")

    def test_large_derived_font_cannot_inherit_a_small_safe_base_indent(self):
        with tempfile.TemporaryDirectory(prefix="indent_inherited_chars_") as directory:
            path = Path(directory) / "inherited.docx"
            document = Document()
            base = document.styles.add_style("SafeTwelvePointIndent", WD_STYLE_TYPE.PARAGRAPH)
            base.font.size = Pt(12)
            _ind(base._element.get_or_add_pPr(), leftChars=200)
            child = document.styles.add_style("LargeDerivedIndent", WD_STYLE_TYPE.PARAGRAPH)
            child.base_style = base
            child.font.size = Pt(36)
            document.add_paragraph("小字号合理缩进", base)
            document.add_paragraph("大字号继承缩进需归零", child)
            document.save(path)
            finalize_docx(path)
            paragraph = _paragraph_by_text(_root(path), "大字号继承缩进需归零")
            self.assertEqual(paragraph.find("w:pPr/w:ind", NS).get(qn("w:leftChars")), "0")
            safe_base = _root(path, "word/styles.xml").xpath(
                ".//w:style[@w:styleId='SafeTwelvePointIndent']/w:pPr/w:ind", namespaces=NS
            )[0]
            self.assertEqual(safe_base.get(qn("w:leftChars")), "200")
            self.assertEqual(finalize_docx(path)["unsafe_indents_reset"], 0)

    def test_unused_large_derived_styles_are_safe_for_future_field_regeneration(self):
        with tempfile.TemporaryDirectory(prefix="indent_future_field_style_") as directory:
            path = Path(directory) / "future_field.docx"
            document = Document()
            base = document.styles.add_style("FutureFieldBase", WD_STYLE_TYPE.PARAGRAPH)
            base.font.size = Pt(12)
            _ind(base._element.get_or_add_pPr(), leftChars=200)
            future = document.styles.add_style("FutureField", WD_STYLE_TYPE.PARAGRAPH)
            future.base_style = base
            future.font.size = Pt(72)
            document.add_paragraph("未来刷新前尚未使用该派生样式")
            document.save(path)
            finalize_docx(path)
            styles = _root(path, "word/styles.xml")
            future_nodes = styles.xpath(
                ".//w:style[@w:styleId='FutureField']/w:pPr/w:ind", namespaces=NS
            )
            self.assertTrue(future_nodes, "Unused derived styles must already be safe before fields create paragraphs")
            self.assertEqual(future_nodes[0].get(qn("w:leftChars")), "0")
            base_node = styles.xpath(
                ".//w:style[@w:styleId='FutureFieldBase']/w:pPr/w:ind", namespaces=NS
            )[0]
            self.assertEqual(base_node.get(qn("w:leftChars")), "200")
            self.assertEqual(finalize_docx(path)["unsafe_indents_reset"], 0)

    def test_complex_script_font_size_also_limits_character_indentation(self):
        with tempfile.TemporaryDirectory(prefix="indent_complex_script_") as directory:
            path = Path(directory) / "complex_script.docx"
            document = Document()
            document.styles["Normal"].font.size = Pt(12)
            paragraph = document.add_paragraph("复杂文字字号不能绕过字符缩进限制")
            paragraph.runs[0].font.size = Pt(12)
            run_cs_size = OxmlElement("w:szCs")
            run_cs_size.set(qn("w:val"), "144")
            paragraph.runs[0]._r.get_or_add_rPr().append(run_cs_size)
            _paragraph_ind(paragraph, leftChars=200)
            style = document.styles.add_style("LargeComplexScriptStyle", WD_STYLE_TYPE.PARAGRAPH)
            style.font.size = Pt(12)
            style_cs_size = OxmlElement("w:szCs")
            style_cs_size.set(qn("w:val"), "144")
            style._element.get_or_add_rPr().append(style_cs_size)
            _ind(style._element.get_or_add_pPr(), rightChars=200)
            document.save(path)
            finalize_docx(path)
            paragraph = _paragraph_by_text(_root(path), "复杂文字字号不能绕过字符缩进限制")
            self.assertEqual(paragraph.find("w:pPr/w:ind", NS).get(qn("w:leftChars")), "0")
            node = _root(path, "word/styles.xml").xpath(
                ".//w:style[@w:styleId='LargeComplexScriptStyle']/w:pPr/w:ind", namespaces=NS
            )[0]
            self.assertEqual(node.get(qn("w:rightChars")), "0")
            self.assertEqual(finalize_docx(path)["unsafe_indents_reset"], 0)

    def test_document_defaults_are_guarded_as_well_as_actual_paragraphs(self):
        with tempfile.TemporaryDirectory(prefix="indent_defaults_") as directory:
            path = Path(directory) / "defaults.docx"
            document = Document()
            document.add_paragraph("使用文档默认缩进的原文")
            styles = document.styles.element
            defaults = styles.find(qn("w:docDefaults"))
            p_default = defaults.find(qn("w:pPrDefault"))
            p_properties = p_default.find(qn("w:pPr"))
            _ind(p_properties, left=9000, rightChars=900)
            document.save(path)
            finalize_docx(path)
            paragraph = _paragraph_by_text(_root(path), "使用文档默认缩进的原文")
            direct = paragraph.find("w:pPr/w:ind", NS)
            self.assertEqual(direct.get(qn("w:left")), "0")
            self.assertEqual(direct.get(qn("w:rightChars")), "0")
            default = _root(path, "word/styles.xml").find("w:docDefaults/w:pPrDefault/w:pPr/w:ind", NS)
            self.assertEqual(default.get(qn("w:left")), "0")
            self.assertEqual(default.get(qn("w:rightChars")), "0")
            self.assertEqual(finalize_docx(path)["unsafe_indents_reset"], 0)

    def test_large_runs_nested_in_simple_fields_or_sdts_bound_character_indents(self):
        with tempfile.TemporaryDirectory(prefix="indent_nested_fonts_") as directory:
            path = Path(directory) / "nested_fonts.docx"
            document = Document()
            document.styles["Normal"].font.size = Pt(12)
            for kind in ("fldSimple", "sdt"):
                paragraph = document.add_paragraph()
                _paragraph_ind(paragraph, leftChars=200)
                container = OxmlElement(f"w:{kind}")
                if kind == "fldSimple":
                    container.set(qn("w:instr"), " PAGE ")
                    run_parent = container
                else:
                    run_parent = OxmlElement("w:sdtContent")
                    container.append(run_parent)
                run = OxmlElement("w:r")
                r_properties = OxmlElement("w:rPr")
                size = OxmlElement("w:sz")
                size.set(qn("w:val"), "144")  # 72 pt within a 12 pt paragraph.
                r_properties.append(size)
                run.append(r_properties)
                text = OxmlElement("w:t")
                text.text = f"nested-{kind}"
                run.append(text)
                run_parent.append(run)
                paragraph._p.append(container)
            document.save(path)
            finalize_docx(path)
            root = _root(path)
            for kind in ("fldSimple", "sdt"):
                with self.subTest(container=kind):
                    paragraph = _paragraph_by_text(root, f"nested-{kind}")
                    self.assertEqual(paragraph.find("w:pPr/w:ind", NS).get(qn("w:leftChars")), "0")
            self.assertEqual(finalize_docx(path)["unsafe_indents_reset"], 0)

    def test_numbering_style_links_resolve_character_indents_at_the_paragraph_font_size(self):
        with tempfile.TemporaryDirectory(prefix="indent_numbering_links_") as directory:
            path = Path(directory) / "numbering_links.docx"
            document = Document()
            document.styles["Normal"].font.size = Pt(12)
            style = document.styles.add_style("LinkedNumberingStyle", WD_STYLE_TYPE.LIST)
            num_pr = OxmlElement("w:numPr")
            style_num_id = OxmlElement("w:numId")
            style_num_id.set(qn("w:val"), "901")
            num_pr.append(style_num_id)
            style._element.get_or_add_pPr().append(num_pr)
            numbering = document.part.numbering_part.element
            linked = OxmlElement("w:abstractNum")
            linked.set(qn("w:abstractNumId"), "900")
            style_link = OxmlElement("w:numStyleLink")
            style_link.set(qn("w:val"), "LinkedNumberingStyle")
            linked.append(style_link)
            numbering.append(linked)
            concrete = OxmlElement("w:abstractNum")
            concrete.set(qn("w:abstractNumId"), "901")
            level = OxmlElement("w:lvl")
            level.set(qn("w:ilvl"), "0")
            p_properties = OxmlElement("w:pPr")
            _ind(p_properties, leftChars=300, hangingChars=200)
            level.append(p_properties)
            concrete.append(level)
            numbering.append(concrete)
            for identity in ("900", "901"):
                num = OxmlElement("w:num")
                num.set(qn("w:numId"), identity)
                abstract_id = OxmlElement("w:abstractNumId")
                abstract_id.set(qn("w:val"), identity)
                num.append(abstract_id)
                numbering.append(num)
            paragraph = document.add_paragraph("编号链接继承的三字缩进")
            paragraph.runs[0].font.size = Pt(48)
            paragraph_num_pr = OxmlElement("w:numPr")
            num_id = OxmlElement("w:numId")
            num_id.set(qn("w:val"), "900")
            paragraph_num_pr.append(num_id)
            paragraph._p.get_or_add_pPr().append(paragraph_num_pr)
            document.save(path)
            finalize_docx(path)
            paragraph = _paragraph_by_text(_root(path), "编号链接继承的三字缩进")
            node = paragraph.find("w:pPr/w:ind", NS)
            self.assertIsNotNone(node, "numStyleLink must not bypass effective-indent protection")
            self.assertEqual(node.get(qn("w:leftChars")), "0")
            self.assertEqual(node.get(qn("w:hangingChars")), "0")
            self.assertEqual(finalize_docx(path)["unsafe_indents_reset"], 0)

    def test_cyclic_numbering_style_links_do_not_crash_or_bypass_direct_guard(self):
        with tempfile.TemporaryDirectory(prefix="indent_numbering_cycle_") as directory:
            path = Path(directory) / "numbering_cycle.docx"
            document = Document()
            style = document.styles.add_style("CyclicNumberingStyle", WD_STYLE_TYPE.LIST)
            num_pr = OxmlElement("w:numPr")
            style_num = OxmlElement("w:numId")
            style_num.set(qn("w:val"), "900")
            num_pr.append(style_num)
            style._element.get_or_add_pPr().append(num_pr)
            numbering = document.part.numbering_part.element
            abstract = OxmlElement("w:abstractNum")
            abstract.set(qn("w:abstractNumId"), "900")
            link = OxmlElement("w:numStyleLink")
            link.set(qn("w:val"), "CyclicNumberingStyle")
            abstract.append(link)
            numbering.append(abstract)
            num = OxmlElement("w:num")
            num.set(qn("w:numId"), "900")
            abstract_id = OxmlElement("w:abstractNumId")
            abstract_id.set(qn("w:val"), "900")
            num.append(abstract_id)
            numbering.append(num)
            paragraph = document.add_paragraph("循环编号不能导致交付失败")
            direct_num_pr = OxmlElement("w:numPr")
            direct_num = OxmlElement("w:numId")
            direct_num.set(qn("w:val"), "900")
            direct_num_pr.append(direct_num)
            paragraph._p.get_or_add_pPr().append(direct_num_pr)
            _paragraph_ind(paragraph, left=9000)
            document.save(path)
            finalize_docx(path)
            paragraph = _paragraph_by_text(_root(path), "循环编号不能导致交付失败")
            self.assertEqual(paragraph.find("w:pPr/w:ind", NS).get(qn("w:left")), "0")
            self.assertEqual(finalize_docx(path)["unsafe_indents_reset"], 0)

    def test_safe_normal_toc_reference_and_rounded_twip_boundary_survive(self):
        with tempfile.TemporaryDirectory(prefix="indent_safe_") as directory:
            path = Path(directory) / "safe.docx"
            document = Document()
            document.styles["Normal"].font.size = Pt(12)
            examples = {
                "正文首行两字": {"firstLineChars": 200},
                "目录四字缩进": {"leftChars": 400},
                "引用悬挂两字": {"leftChars": 200, "hangingChars": 200},
                "左右厘米边界": {"left": Cm(2).twips, "right": Cm(2).twips},
                "逻辑字符边界": {"startChars": 400, "endChars": 400},
                "首行厘米边界": {"firstLine": Cm(2).twips},
                "悬挂厘米边界": {"hanging": Cm(2).twips},
            }
            for text, attributes in examples.items():
                paragraph = document.add_paragraph(text)
                paragraph.runs[0].font.size = Pt(12)
                _paragraph_ind(paragraph, **attributes)
            document.save(path)
            finalize_docx(path)
            root = _root(path)
            for text, attributes in examples.items():
                node = _paragraph_by_text(root, text).find("w:pPr/w:ind", NS)
                for name, value in attributes.items():
                    self.assertEqual(node.get(qn(f"w:{name}")), str(value), (text, name))

    def test_table_indents_reject_bad_dxa_but_keep_safe_offsets(self):
        with tempfile.TemporaryDirectory(prefix="indent_tables_") as directory:
            path = Path(directory) / "tables.docx"
            document = Document()
            values = ("-1", "NaN", "INF", "bad", "99999", "1135", "1134", "720", "0")
            for value in values:
                table = document.add_table(1, 1)
                table.cell(0, 0).text = f"偏移 {value}"
                indent = OxmlElement("w:tblInd")
                indent.set(qn("w:type"), "dxa")
                indent.set(qn("w:w"), value)
                table._tbl.tblPr.append(indent)
            document.save(path)
            finalize_docx(path)
            nodes = _root(path).xpath(".//w:tblInd", namespaces=NS)
            self.assertEqual([node.get(qn("w:w")) for node in nodes], ["0"] * 6 + ["1134", "720", "0"])

    def test_idempotence_preserves_text_image_fields_bookmarks_and_package_parts(self):
        with tempfile.TemporaryDirectory(prefix="indent_integrity_") as directory:
            path = Path(directory) / "integrity.docx"
            document = Document()
            paragraph = document.add_paragraph("作者原文、符号 αβ 与图片必须保留。")
            _paragraph_ind(paragraph, left=8000, firstLineChars=200)
            paragraph.add_run().add_picture(BytesIO(PIXEL_PNG), width=Cm(0.2))
            bookmark = OxmlElement("w:bookmarkStart")
            bookmark.set(qn("w:id"), "42")
            bookmark.set(qn("w:name"), "valid_target")
            paragraph._p.append(bookmark)
            paragraph.add_run("引用目标")
            end = OxmlElement("w:bookmarkEnd")
            end.set(qn("w:id"), "42")
            paragraph._p.append(end)
            begin = OxmlElement("w:fldChar")
            begin.set(qn("w:fldCharType"), "begin")
            paragraph.add_run()._r.append(begin)
            instruction = OxmlElement("w:instrText")
            instruction.text = " REF valid_target \\h "
            paragraph.add_run()._r.append(instruction)
            separate = OxmlElement("w:fldChar")
            separate.set(qn("w:fldCharType"), "separate")
            paragraph.add_run()._r.append(separate)
            paragraph.add_run("引用目标")
            field_end = OxmlElement("w:fldChar")
            field_end.set(qn("w:fldCharType"), "end")
            paragraph.add_run()._r.append(field_end)
            simple = OxmlElement("w:fldSimple")
            simple.set(qn("w:instr"), " PAGE ")
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = "1"
            run.append(text)
            simple.append(run)
            paragraph._p.append(simple)
            document.save(path)
            before = _package(path)
            self.assertGreater(finalize_docx(path)["unsafe_indents_reset"], 0)
            once = _package(path)
            self.assertEqual(set(once), set(before))
            self.assertEqual(_semantic_snapshot(once), _semantic_snapshot(before))
            self.assertEqual(len(Document(path).inline_shapes), 1)
            self.assertEqual(finalize_docx(path)["unsafe_indents_reset"], 0)
            self.assertEqual(_package(path), once)


class ProcessorFinalIndentGuardTests(unittest.TestCase):
    def test_final_word_refresh_cannot_reintroduce_bad_cover_header_or_style_indents(self):
        with tempfile.TemporaryDirectory(prefix="indent_last_refresh_") as directory:
            source = Path(directory) / "original.docx"
            output = Path(directory) / "delivery.docx"
            document = Document()
            document.add_paragraph("大学毕业论文封面作者信息")
            document.add_paragraph("独创性声明原文应完整保留")
            document.add_paragraph("第一章 绪论", style="Heading 1")
            document.add_paragraph("作者正文不得丢失。")
            document.sections[0].header.paragraphs[0].text = "作者页眉"
            document.save(source)
            original_bytes = source.read_bytes()
            refresh_calls = []

            def refresh(path):
                refresh_calls.append(Path(path))
                if len(refresh_calls) != 2:
                    return
                refreshed = Document(path)
                cover = next(p for p in refreshed.paragraphs if p.text == "大学毕业论文封面作者信息")
                _paragraph_ind(cover, left=9000, endChars=900)
                _paragraph_ind(refreshed.sections[0].header.paragraphs[0], start=9000)
                style = refreshed.styles.add_style("WordRefreshUnsafeStyle", WD_STYLE_TYPE.PARAGRAPH)
                _ind(style._element.get_or_add_pPr(), right=9000)
                for paragraph in refreshed.paragraphs:
                    if paragraph.style.name.casefold().replace(" ", "") == "toc1":
                        _paragraph_ind(paragraph, left=9000)
                refreshed.save(path)

            with patch("word_formatter.core.processor.os", SimpleNamespace(name="nt")), patch(
                "word_formatter.core.processor.WordDocumentConverter.update_fields_in_place", side_effect=refresh
            ):
                DocumentProcessor().process(source, DocumentRules(), output)
            self.assertEqual(len(refresh_calls), 2)
            self.assertEqual(source.read_bytes(), original_bytes)
            root = _root(output)
            cover_ind = _paragraph_by_text(root, "大学毕业论文封面作者信息").find("w:pPr/w:ind", NS)
            self.assertEqual(cover_ind.get(qn("w:left")), "0")
            self.assertEqual(cover_ind.get(qn("w:endChars")), "0")
            header = _root(output, "word/header1.xml")
            self.assertEqual(header.find(".//w:ind", NS).get(qn("w:start")), "0")
            style = _root(output, "word/styles.xml").xpath(
                ".//w:style[@w:styleId='WordRefreshUnsafeStyle']/w:pPr/w:ind", namespaces=NS
            )[0]
            self.assertEqual(style.get(qn("w:right")), "0")
            self.assertIn("作者正文不得丢失。", root.xpath(".//w:t/text()", namespaces=NS))
            self.assertIn("独创性声明原文应完整保留", root.xpath(".//w:t/text()", namespaces=NS))
            self.assertEqual(finalize_docx(output)["unsafe_indents_reset"], 0)


if __name__ == "__main__":
    unittest.main()
