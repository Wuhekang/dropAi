from argparse import Namespace
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement

from format_cli import CliInputError, _restore_analyzed_rules, run_job
from word_formatter.models.rules import DocumentRules


class ConfirmedExecutionTests(unittest.TestCase):
    def execute(self, envelope, instructions="正文小二号。", analyze_only=False):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        root = Path(directory.name)
        source, template, output = (root / name for name in ("source.docx", "template.docx", "out.docx"))
        doc = Document()
        # The production crash happened before processing a plain lxml SDT.
        sdt, content = OxmlElement("w:sdt"), OxmlElement("w:sdtContent")
        p, run, text = OxmlElement("w:p"), OxmlElement("w:r"), OxmlElement("w:t")
        text.text = "原稿内容控件文字必须保留"
        run.append(text); p.append(run); content.append(p); sdt.append(content)
        doc.element.body.insert(0, sdt)
        doc.add_paragraph("第一章 绪论", style="Heading 1")
        doc.add_paragraph("这是作者正文，必须保持文字内容不变，只修改格式。")
        doc.add_paragraph("1.1 二级标题", style="Heading 2")
        doc.add_paragraph("图1-1 系统结构")
        doc.add_paragraph("参考文献", style="Heading 1")
        doc.add_paragraph("[1] 文献内容。")
        doc.save(source)
        doc = Document()
        doc.add_paragraph("模板没有明确提供各级标题和参考文献格式。")
        doc.save(template)
        source_bytes = source.read_bytes()
        instruction_path, rules_path = root / "instructions.txt", root / "rules.json"
        instruction_path.write_text(instructions, encoding="utf-8")
        rules_path.write_text(json.dumps(envelope, ensure_ascii=False), encoding="utf-8")
        args = Namespace(source=str(source), template=str(template), output=str(output),
                         result_json=str(root / "result.json"), instructions_file=str(instruction_path),
                         rules_file=str(rules_path), use_doubao=True, analyze_only=analyze_only)
        with patch("format_cli.DoubaoRuleParser", side_effect=AssertionError("AI must never run while formatting")), \
             patch("format_cli.read_template_text", side_effect=AssertionError("confirmed plans must not be re-inferred")), \
             patch("format_cli.TemplateRuleExtractor", side_effect=AssertionError("confirmed plans must not be re-inferred")), \
             patch("word_formatter.core.processor.WordDocumentConverter.update_fields_in_place"):
            payload = run_job(args)
        self.assertTrue(payload["success"])
        self.assertEqual(source.read_bytes(), source_bytes)
        if not analyze_only:
            self.assertTrue(output.is_file())
            delivered = Document(output)
            self.assertIn("原稿内容控件文字必须保留", "".join(delivered.element.itertext()))
            body = next(p for p in delivered.paragraphs if p.text.startswith("这是作者正文"))
            self.assertTrue(all(run.font.size.pt == payload["ruleSummary"]["normalText"]["fontSizePt"] for run in body.runs if run.text))
        return payload

    def test_legacy_confirmation_uses_defaults_and_edits_without_ai_or_template_inference(self):
        payload = self.execute({"confirmationVersion": 2, "editableRules": {
            "body": {"normal": {"fontSizePt": 14, "bold": False, "leftIndentCm": 0.5}}
        }})
        rules = payload["ruleSummary"]
        self.assertEqual(rules["normalText"]["fontSizePt"], 14)
        self.assertEqual(rules["normalText"]["leftIndentCm"], 0.5)
        self.assertEqual(rules["heading1"]["fontSizePt"], 16)
        self.assertEqual(rules["figureCaption"]["fontSizePt"], 10.5)
        for key in ("normalText", "heading1", "heading2", "heading3", "tocTitle", "toc1", "toc2", "toc3", "figureCaption", "tableCaption", "reference", "table"):
            self.assertTrue(rules[key]["enabled"], key)
        self.assertFalse(payload["templateAnalysis"]["copyFrontMatter"])

    def test_current_confirmation_values_are_authoritative_and_partial_snapshots_fill_defaults(self):
        payload = self.execute({"confirmationVersion": 2,
            "analyzedRules": {"normal_text": {"font_size_pt": 14, "bold": True, "alignment": "left"}},
            "editableRules": {"body": {"normal": {"fontSizePt": 12, "bold": True, "alignment": "left", "latinFont": "Arial"}}},
            "templateAnalysis": {"copyFrontMatter": False, "documentKind": "specification"}
        })
        body = payload["ruleSummary"]["normalText"]
        self.assertEqual(body["fontSizePt"], 12)
        self.assertTrue(body["bold"])
        self.assertEqual(body["alignment"], "left")
        self.assertEqual(body["latinFont"], "Arial")
        self.assertEqual(payload["ruleSummary"]["heading1"]["chineseFont"], "黑体")

    def test_empty_confirmed_form_executes_defaults(self):
        payload = self.execute({"confirmationVersion": 2, "editableRules": {}}, instructions="")
        self.assertEqual(payload["ruleSummary"]["normalText"]["fontSizePt"], 12)
        self.assertTrue(payload["ruleSummary"]["heading3"]["enabled"])

    def test_missing_recognition_is_returned_as_editable_enabled_defaults(self):
        payload = self.execute({"editableRules": {}}, instructions="", analyze_only=True)
        self.assertTrue(payload["analysisReady"])
        self.assertTrue(payload["editableRules"]["headings"]["level2"]["enabled"])
        self.assertEqual(payload["editableRules"]["body"]["normal"]["latinFont"], "Times New Roman")
        self.assertEqual(payload["templateAnalysis"]["ruleEvidence"]["heading_2"]["status"], "unconfirmed")

    def test_changed_template_is_still_rejected_instead_of_silently_using_wrong_cover(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "rules.json"
            path.write_text(json.dumps({"templateSha256": "different", "analyzedRules": DocumentRules().to_dict()}), encoding="utf-8")
            with self.assertRaisesRegex(CliInputError, "模板已变化"):
                _restore_analyzed_rules(path, "current")


if __name__ == "__main__":
    unittest.main()
