from __future__ import annotations

from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.opc.constants import CONTENT_TYPE as CT, RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from lxml import etree

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from word_formatter.core.finalizer import finalize_docx
from word_formatter.core.processor import DocumentProcessor
from word_formatter.core.xml_utils import OOXML_NAMESPACES, xpath
from word_formatter.models.rules import DocumentRules


W_NS = OOXML_NAMESPACES["w"]
M_NS = OOXML_NAMESPACES["m"]


def _sdt(prefix="w"):
    # etree.fromstring deliberately produces plain lxml nodes, not CT_* nodes.
    # The TOC has a misleading cached heading that must not become the body.
    return etree.fromstring(f'''<s:sdt xmlns:s="{W_NS}" xmlns:{prefix}="{W_NS}">
      <s:sdtContent><s:p><s:r><s:fldChar s:fldCharType="begin"/></s:r>
        <s:r><s:instrText> TOC \\o "1-3" \\h </s:instrText></s:r>
        <s:r><s:fldChar s:fldCharType="separate"/></s:r>
        <s:r><s:t>第一章 目录缓存而非作者正文</s:t></s:r>
        <s:r><s:fldChar s:fldCharType="end"/></s:r>
      </s:p></s:sdtContent></s:sdt>'''.encode())


def _change_document_prefix(path, prefix):
    """Change only namespace spelling, keeping names and document content."""
    with ZipFile(path) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    data = parts["word/document.xml"]
    if prefix == "default":
        data = data.replace(b'xmlns:w=', b'xmlns=')
        data = data.replace(b'<w:', b'<').replace(b'</w:', b'</')
        data = data.replace(b' w:', b' attr:')
        data = data.replace(b'<document ', f'<document xmlns:attr="{W_NS}" '.encode(), 1)
    else:
        data = data.replace(b'xmlns:w=', f'xmlns:{prefix}='.encode())
        data = data.replace(b'w:', f'{prefix}:'.encode())
    parts["word/document.xml"] = data
    with ZipFile(path, "w", ZIP_DEFLATED) as package:
        for name, blob in parts.items():
            package.writestr(name, blob)


class NamespaceQueryTests(unittest.TestCase):
    def test_plain_lxml_and_python_docx_return_original_nodes(self):
        typed = OxmlElement("w:p")
        typed.append(OxmlElement("w:r"))
        plain = etree.fromstring(f'<p xmlns="{W_NS}"><r/></p>'.encode())
        self.assertIsInstance(typed, BaseOxmlElement)
        self.assertNotIsInstance(plain, BaseOxmlElement)
        for node in (typed, plain):
            with self.subTest(type=type(node).__name__):
                result = xpath(node, "./w:r")
                self.assertEqual(len(result), 1)
                self.assertIs(result[0], node[0])
                result[0].set(qn("w:rsidR"), "00000001")
                self.assertEqual(node[0].get(qn("w:rsidR")), "00000001")

    def test_queries_ignore_document_prefix_spelling_and_rebinding(self):
        roots = [
            f'<w:p xmlns:w="{W_NS}"><w:r><w:t>text</w:t></w:r></w:p>',
            f'<word:p xmlns:word="{W_NS}"><word:r><word:t>text</word:t></word:r></word:p>',
            f'<p xmlns="{W_NS}"><r><t>text</t></r></p>',
            f'<word:p xmlns:word="{W_NS}" xmlns:w="urn:not-word"><word:r><word:t>text</word:t></word:r><w:t>decoy</w:t></word:p>',
        ]
        for xml in roots:
            with self.subTest(xml=xml):
                self.assertEqual(xpath(etree.fromstring(xml.encode()), ".//w:t/text()"), ["text"])
        math = etree.fromstring(f'<math:oMathPara xmlns:math="{M_NS}"><math:oMath/></math:oMathPara>'.encode())
        self.assertIs(xpath(math, ".//m:oMath")[0], math[0])

    def test_genuine_xpath_errors_are_not_silenced(self):
        node = OxmlElement("w:p")
        with self.assertRaises(etree.XPathEvalError):
            xpath(node, ".//unknown:element")
        with self.assertRaises(etree.XPathEvalError):
            xpath(node, "[")


class ProcessorNamespaceTests(unittest.TestCase):
    def test_main_and_retained_start_scan_plain_sdt_altchunk_and_math(self):
        for spelling in ("w", "word"):
            with self.subTest(spelling=spelling):
                document = Document()
                document.add_paragraph("论文封面")
                document.add_paragraph("独创性声明")
                document.add_heading("第一章 绪论", 1)
                for node in (
                    _sdt(spelling),
                    OxmlElement("w:altChunk"),
                    OxmlElement("m:oMathPara"),
                ):
                    self.assertNotIsInstance(node, BaseOxmlElement)
                    document.element.body.insert(2, node)
                self.assertEqual(DocumentProcessor._main_content_start(document), 3)
                self.assertEqual(DocumentProcessor._retained_content_start(document, 3), 3)

    def test_review_cleanup_handles_plain_lxml_descendants_without_dropping_author_text(self):
        document = Document()
        document.add_paragraph("论文封面")
        document.add_heading("第一章 绪论", 1)
        sdt = etree.fromstring(f'''<s:sdt xmlns:s="{W_NS}"><s:sdtContent><s:p>
          <s:r><s:rPr><s:color s:val="FF0000"/></s:rPr><s:t>请删除此提示</s:t></s:r>
          <s:r><s:t>作者信息应保留</s:t></s:r>
        </s:p></s:sdtContent></s:sdt>'''.encode())
        document.element.body.insert(1, sdt)
        self.assertEqual(DocumentProcessor._remove_template_review_artifacts(document, 2), 1)
        self.assertEqual(xpath(sdt, ".//w:t/text()"), ["作者信息应保留"])

    def test_real_docx_process_preserves_front_toc_math_altchunk_and_generic_parts(self):
        for prefix in ("w", "word", "default"):
            with self.subTest(prefix=prefix), tempfile.TemporaryDirectory(prefix="namespace_docx_") as directory:
                source = Path(directory) / "source.docx"
                output = Path(directory) / "output.docx"
                document = Document()
                cover = document.add_paragraph("论文封面作者姓名")
                ind = OxmlElement("w:ind")
                ind.set(qn("w:left"), "9000")
                cover._p.get_or_add_pPr().append(ind)
                document.add_paragraph("独创性声明原文")
                document.add_heading("第一章 绪论", 1)
                document.add_paragraph("作者正文必须完整保留。")
                document.element.body.insert(2, _sdt())
                formula = OxmlElement("m:oMathPara")
                math = OxmlElement("m:oMath")
                run = OxmlElement("m:r")
                text = OxmlElement("m:t")
                text.text = "x+y=1"
                run.append(text)
                math.append(run)
                formula.append(math)
                document.element.body.insert(5, formula)

                html = Part(PackURI("/word/embedded.html"), "text/html", b"<html><body>Imported text</body></html>", document.part.package)
                chunk = OxmlElement("w:altChunk")
                chunk.set(qn("r:id"), document.part.relate_to(html, RT.A_F_CHUNK))
                document.element.body.insert(6, chunk)
                footnotes_xml = f'''<notes:footnotes xmlns:notes="{W_NS}">
                  <notes:footnote notes:id="1"><notes:p><notes:pPr><notes:ind notes:left="9000"/></notes:pPr>
                    <notes:r><notes:t>脚注原文</notes:t></notes:r></notes:p></notes:footnote>
                </notes:footnotes>'''.encode()
                footnotes = Part(PackURI("/word/footnotes.xml"), CT.WML_FOOTNOTES, footnotes_xml, document.part.package)
                document.part.relate_to(footnotes, RT.FOOTNOTES)
                footnote_ref = OxmlElement("w:footnoteReference")
                footnote_ref.set(qn("w:id"), "1")
                document.paragraphs[-1].add_run()._r.append(footnote_ref)
                document.save(source)
                if prefix != "w":
                    _change_document_prefix(source, prefix)
                original = source.read_bytes()

                loaded = Document(source)
                self.assertEqual(DocumentProcessor._main_content_start(loaded), 3)
                self.assertIs(type(loaded.part.part_related_by(RT.FOOTNOTES)), Part)
                if prefix != "w":
                    self.assertNotIn("w", loaded.element.nsmap)
                with patch("word_formatter.core.processor.os", SimpleNamespace(name="posix")):
                    DocumentProcessor().process(source, DocumentRules(), output)
                self.assertEqual(source.read_bytes(), original)
                self.assertTrue(output.with_suffix(".log.json").is_file())
                self.assertFalse(output.with_suffix(".failed.log.json").exists())
                with ZipFile(output) as package:
                    root = etree.fromstring(package.read("word/document.xml"))
                    footnotes_root = etree.fromstring(package.read("word/footnotes.xml"))
                    self.assertEqual(package.read("word/embedded.html"), html.blob)
                texts = xpath(root, ".//w:t/text()")
                for text in ("论文封面作者姓名", "独创性声明原文", "第一章 绪论", "作者正文必须完整保留。", "第一章 目录缓存而非作者正文"):
                    self.assertIn(text, texts)
                self.assertEqual(len(xpath(root, ".//w:sdt")), 1)
                self.assertEqual(len(xpath(root, ".//w:altChunk")), 1)
                self.assertEqual(len(xpath(root, ".//w:instrText[contains(., 'TOC ')]")), 1)
                self.assertEqual(xpath(root, ".//m:t/text()"), ["x+y=1"])
                self.assertEqual(xpath(footnotes_root, ".//w:t/text()"), ["脚注原文"])
                self.assertEqual(xpath(footnotes_root, ".//w:ind/@w:left"), ["0"])
                cover_xml = xpath(root, ".//w:p[w:r/w:t='论文封面作者姓名']")[0]
                self.assertEqual(xpath(cover_xml, "./w:pPr/w:ind/@w:left"), ["0"])
                self.assertEqual(finalize_docx(output)["unsafe_indents_reset"], 0)
                self.assertIn("作者正文必须完整保留。", [p.text for p in Document(output).paragraphs])


if __name__ == "__main__":
    unittest.main()
