from __future__ import annotations

from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from word_formatter.core.rule_parser import extract_explicit_template_fields
from word_formatter.core.template_extractor import TemplateRuleExtractor
from word_formatter.core.template_text import read_template_text


def block(text, identity="b1", region="main"):
    return {"id": identity, "kind": "paragraph", "text": text, "semanticRegion": region}


class WrittenPriorityTests(unittest.TestCase):
    def test_inline_written_size_wins_over_incorrect_sample_style(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "template.docx"
            doc = Document()
            title = doc.add_paragraph("目 录（三号黑体字）")
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph("前言…………1").runs[0].font.size = Pt(14)
            doc.add_paragraph("1、绪论…………2").runs[0].font.size = Pt(14)
            doc.add_paragraph("1.1 目的…………2").runs[0].font.size = Pt(14)
            title = doc.add_paragraph("前 言（三号黑体字）")
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for prefix in ("1 ", "1.1 ", "1.1.1 "):
                p = doc.add_paragraph(prefix + "XXXXXXXX（三号黑体字）")
                p.runs[0].font.size = Pt(10.5)
            doc.add_paragraph("XXXXXXX（小四号宋体字）").runs[0].font.size = Pt(16)
            doc.save(path)
            rules = TemplateRuleExtractor().extract(path).rules
            for role in ("heading_1", "heading_2", "heading_3", "toc_title"):
                rule = getattr(rules, role)
                self.assertEqual((rule.chinese_font, rule.font_size_pt), ("黑体", 16))
            for role in ("toc_1", "toc_2", "toc_3"):
                self.assertEqual(getattr(rules, role).font_size_pt, 14)
            self.assertEqual(rules.toc_title.alignment, "center")
            self.assertEqual(rules.normal_text.font_size_pt, 12)
            self.assertIn("heading_2", read_template_text(path)["explicitRuleFields"])

    def test_toc_heading_labels_do_not_override_body_heading_rules(self):
        result = extract_explicit_template_fields([
            block("一级标题：四号、宋体、顶格\n二级标题：小四、宋体、首行缩进2字符", "toc", "toc"),
            block("正文一级标题\n三号、宋体、加粗、居中\n二级标题：宋体、加粗、四号，1.5倍行距", "body"),
        ])
        self.assertEqual(result["toc_1"]["rule"]["font_size_pt"], 14)
        self.assertEqual(result["toc_2"]["rule"]["font_size_pt"], 12)
        self.assertEqual(result["heading_1"]["rule"]["font_size_pt"], 16)
        self.assertEqual(result["heading_2"]["rule"]["font_size_pt"], 14)
        self.assertEqual(result["heading_2"]["fieldEvidence"]["font_size_pt"], ["body"])

    def test_toc_without_preface_stops_at_actual_first_chapter(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "chapter.docx"
            doc = Document()
            doc.add_paragraph("2024 年 5 月").runs[0].font.size = Pt(26)
            doc.add_paragraph("目录")
            doc.add_paragraph("第一章 绪论…………1").runs[0].font.size = Pt(14)
            heading = doc.add_heading("第一章 绪论", 1)
            heading.runs[0].font.size = Pt(16)
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph("这是一段正文，用于检验目录范围之后仍然可以选取正常的正文格式样例。").runs[0].font.size = Pt(12)
            doc.save(path)
            rules = TemplateRuleExtractor().extract(path).rules
            self.assertEqual(rules.heading_1.font_size_pt, 16)
            self.assertEqual(rules.heading_1.alignment, "center")
            self.assertEqual(rules.normal_text.font_size_pt, 12)

    def test_reference_body_and_table_contents_have_independent_scopes(self):
        result = extract_explicit_template_fields([
            block("正文：宋体、小四、不加粗", "body"),
            block("正文\n单倍行距、中文：五号、宋体", "ref", "reference"),
            block("表名：黑体四号、居中；表内：五号、中文宋体，Times New Roman字体", "table"),
        ])
        self.assertEqual(result["normal_text"]["rule"]["font_size_pt"], 12)
        self.assertFalse(result["normal_text"]["rule"]["bold"])
        self.assertEqual(result["reference"]["rule"]["font_size_pt"], 10.5)
        self.assertEqual(result["table_caption"]["rule"]["font_size_pt"], 14)
        self.assertEqual(result["table"]["rule"]["font_size_pt"], 10.5)
        self.assertEqual(result["table"]["rule"]["chinese_font"], "宋体")

    def test_global_font_cannot_erase_specific_heading_font(self):
        result = extract_explicit_template_fields([
            block("全文：宋体、小四", "global"), block("一级标题：黑体、三号", "specific")])
        self.assertEqual(result["heading_1"]["rule"]["chinese_font"], "黑体")
        self.assertEqual(result["heading_1"]["rule"]["font_size_pt"], 16)
        self.assertEqual(result["normal_text"]["rule"]["font_size_pt"], 12)

    def test_conflicting_specific_requirements_are_not_silently_overwritten(self):
        result = extract_explicit_template_fields([
            block("二级标题：四号", "a"), block("二级标题：小四", "b")])
        self.assertNotIn("font_size_pt", result.get("heading_2", {}).get("rule", {}))

    def test_global_body_layout_is_not_promoted_to_all_headings(self):
        result = extract_explicit_template_fields([block("全文：宋体、小四、加粗、居中、首行缩进2字符")])
        self.assertEqual(result["normal_text"]["rule"]["font_size_pt"], 12)
        self.assertEqual(result["heading_2"]["rule"], {"chinese_font": "宋体"})

    def test_inline_toc_level_requirement_does_not_leak_to_main_heading(self):
        for text in ("目录二级小四", "目录二级标题：小四", "目录：二级标题小四"):
            with self.subTest(text=text):
                result = extract_explicit_template_fields([block(text)])
                self.assertEqual(result["toc_2"]["rule"]["font_size_pt"], 12)
                self.assertNotIn("heading_2", result)

    def test_pdf_page_margin_labels_are_parsed_with_units(self):
        result = extract_explicit_template_fields([block("A4纸\n页面设置\n上：2.5 cm\n下：2.5cm\n左：3.0cm\n右：30mm")])
        self.assertEqual(result["page_setup"]["rule"]["margin_left_mm"], 30)
        self.assertEqual(result["page_setup"]["rule"]["margin_top_mm"], 25)
        self.assertEqual(result["page_setup"]["rule"]["paper"], "A4")

    def test_standalone_toc_spacing_only_applies_with_verified_scope(self):
        result = extract_explicit_template_fields([block("目录\n1.5 倍行距\n二级标题：小四、单倍行距", region="toc")])
        self.assertEqual(result["toc_1"]["rule"]["line_spacing_mode"], "1.5")
        self.assertEqual(result["toc_2"]["rule"]["line_spacing_mode"], "single")
        self.assertEqual(result["toc_3"]["rule"]["line_spacing_mode"], "1.5")
        self.assertEqual(extract_explicit_template_fields([block("1.5 倍行距")]), {})


if __name__ == "__main__":
    unittest.main()
