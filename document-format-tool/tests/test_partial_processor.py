from __future__ import annotations

import base64
from io import BytesIO
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

TOOL_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(TOOL_ROOT))

from word_formatter.core.processor import DocumentProcessor
from word_formatter.core.word_converter import WordDocumentConverter
from word_formatter.models.results import ChangeRecord, ProcessResult
from word_formatter.models.rules import DocumentRules


class PartialProcessorTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix="partial-processor-test-")
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.source = self.root / "source.docx"
        self.output = self.root / "output.docx"
        self.rules = DocumentRules()
        self.rules.normal_text.font_size_pt = 14
        self.rules.normal_text.font_size_name = "四号"
        self.rules.normal_text.alignment = "justify"
        self.rules.page_setup.margin_top_mm = 28
        self.refresh = patch.object(WordDocumentConverter, "update_fields_in_place").start()
        self.addCleanup(patch.stopall)
        patch("word_formatter.core.processor.time.sleep").start()

    def make_source(self, cover=False, heading=True):
        document = Document()
        if cover:
            document.add_paragraph("毕业论文封面").add_run("原版封面").font.size = Pt(30)
        if heading:
            document.add_heading("第一章 绪论", level=1)
        bad = document.add_paragraph("坏正文段落")
        bad.runs[0].font.size = Pt(9)
        document.add_paragraph("健康的普通正文内容，应当继续真实应用指定字号。")
        field = bad.add_run()
        for kind in ("begin", "separate", "end"):
            marker = OxmlElement("w:fldChar")
            marker.set(qn("w:fldCharType"), kind)
            field._r.append(marker)
            if kind == "begin":
                instruction = OxmlElement("w:instrText")
                instruction.text = " PAGE "
                field._r.append(instruction)
        image = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jG3cAAAAASUVORK5CYII=")
        bad.add_run().add_picture(BytesIO(image))
        for name in ("坏表", "健康表"):
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = name
            table.cell(1, 1).text = "原始数据123"
        document.save(self.source)
        self.original = self.source.read_bytes()

    def run_processor(self, **kwargs):
        result = DocumentProcessor().process(self.source, self.rules, self.output, **kwargs)
        self.assertEqual(self.source.read_bytes(), self.original)
        return result, Document(self.output)

    def assert_preserved(self, output):
        texts = [p.text for p in output.paragraphs]
        self.assertIn("坏正文段落", texts)
        self.assertIn("健康的普通正文内容，应当继续真实应用指定字号。", texts)
        with zipfile.ZipFile(self.source) as source, zipfile.ZipFile(self.output) as formatted:
            media = [name for name in source.namelist() if name.startswith("word/media/")]
            self.assertTrue(media)
            for name in media:
                self.assertEqual(source.read(name), formatted.read(name))
        self.assertTrue(any(" PAGE " in (node.text or "") for node in output.element.iter(qn("w:instrText"))))

    def assert_healthy_body_formatted(self, document):
        paragraph = next(p for p in document.paragraphs if p.text.startswith("健康的普通正文"))
        self.assertEqual(paragraph.runs[0].font.size.pt, 14)
        self.assertEqual(paragraph.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)

    def test_bad_paragraph_rolls_back_and_later_body_is_formatted(self):
        self.make_source()
        original = DocumentProcessor._format_paragraph

        def fail_one(paragraph, rule):
            if paragraph.text == "坏正文段落":
                paragraph.text = "MUST BE ROLLED BACK"
                raise RuntimeError("Undefined namespace prefix: isolated paragraph")
            return original(paragraph, rule)

        with patch.object(DocumentProcessor, "_format_paragraph", side_effect=fail_one):
            result, formatted = self.run_processor()
        self.assert_healthy_body_formatted(formatted)
        bad = next(p for p in formatted.paragraphs if p.text == "坏正文段落")
        self.assertEqual(bad.runs[0].font.size.pt, 9)
        self.assert_preserved(formatted)
        skipped = [r for r in result.records if r.status == "skipped" and r.item == "普通正文"]
        self.assertEqual(len(skipped), 1)
        self.assertFalse(any(r.status == "success" and r.item == "普通正文" and r.paragraph_index == skipped[0].paragraph_index for r in result.records))
        self.assertTrue(result.partial_success)
        log = json.loads(self.output.with_suffix(".log.json").read_text(encoding="utf-8"))
        self.assertTrue(log["partial_success"])
        self.assertEqual(log["skipped_count"], result.skipped_count)
        self.assertEqual(log["changed_count"], result.changed_count)

    def test_failed_toc_stage_rolls_back_its_changes_and_false_records(self):
        self.make_source()

        def fail_toc(document, rules, result, content_start):
            document.paragraphs[0].text = "CORRUPTED"
            result.records.append(ChangeRecord(None, "false-success", "", "", ""))
            raise RuntimeError("Undefined namespace prefix")

        with patch.object(DocumentProcessor, "_ensure_toc", side_effect=fail_toc):
            result, formatted = self.run_processor()
        self.assert_healthy_body_formatted(formatted)
        self.assert_preserved(formatted)
        self.assertIn("第一章 绪论", [p.text for p in formatted.paragraphs])
        self.assertFalse(any(r.item == "false-success" for r in result.records))
        self.assertTrue(any(r.item == "自动目录" and r.status == "skipped" for r in result.records))
        self.assertAlmostEqual(formatted.sections[-1].top_margin.mm, 28, places=1)

    def test_late_failed_module_keeps_earlier_body_and_later_tables(self):
        self.make_source()

        def fail_headings(document, rules, result, start):
            next(p for p in document.paragraphs if p.text.startswith("健康的普通正文")).text = "CORRUPTED"
            raise RuntimeError("bad heading metadata")

        with patch.object(DocumentProcessor, "_apply_headings", side_effect=fail_headings):
            result, formatted = self.run_processor()
        self.assert_healthy_body_formatted(formatted)
        self.assert_preserved(formatted)
        self.assertIsNotNone(formatted.tables[-1]._tbl.tblPr.find(qn("w:tblBorders")))
        self.assertTrue(any(r.item == "标题" and r.status == "skipped" for r in result.records))

    def test_bad_table_does_not_skip_later_table(self):
        self.make_source()
        original = DocumentProcessor._set_table_borders

        def fail_one(table, rule):
            if table.cell(0, 0).text == "坏表":
                table.cell(0, 0).text = "CORRUPTED"
                raise ValueError("bad table metadata")
            return original(table, rule)

        with patch.object(DocumentProcessor, "_set_table_borders", side_effect=fail_one):
            result, formatted = self.run_processor()
        self.assertEqual(formatted.tables[0].cell(0, 0).text, "坏表")
        self.assertIsNone(formatted.tables[0]._tbl.tblPr.find(qn("w:tblBorders")))
        self.assertIsNotNone(formatted.tables[1]._tbl.tblPr.find(qn("w:tblBorders")))
        self.assertTrue(any(r.item == "第 1 个表格" and r.status == "skipped" for r in result.records))
        self.assertTrue(any(r.item == "第 2 个表格" and r.status == "success" for r in result.records))
        self.assert_healthy_body_formatted(formatted)

    def test_cover_copy_failure_uses_original_and_formats_body(self):
        self.make_source(cover=True)
        with patch.object(DocumentProcessor, "_compose_with_template_front", side_effect=RuntimeError("broken template")):
            result, formatted = self.run_processor(template_path=self.root / "template.docx")
        self.assertEqual(formatted.paragraphs[0].runs[-1].font.size.pt, 30)
        self.assert_healthy_body_formatted(formatted)
        self.assertTrue(any(r.item == "模板前置内容" and r.status == "skipped" for r in result.records))

    def test_failed_content_recognition_preserves_cover_and_finds_real_heading(self):
        self.make_source(cover=True)
        with patch.object(DocumentProcessor, "_main_content_start", side_effect=RuntimeError("Undefined namespace prefix")):
            result, formatted = self.run_processor()
        self.assertEqual(formatted.paragraphs[0].runs[-1].font.size.pt, 30)
        self.assert_healthy_body_formatted(formatted)
        self.assertTrue(any(r.item == "正文起点识别" and r.status == "skipped" for r in result.records))

    def test_uncertain_content_start_does_not_apply_body_rules_to_cover(self):
        self.make_source(cover=True, heading=False)
        with patch.object(DocumentProcessor, "_main_content_start", side_effect=RuntimeError("broken analyzer")):
            result, formatted = self.run_processor()
        self.assertEqual(formatted.paragraphs[0].runs[-1].font.size.pt, 30)
        self.assertFalse(any(r.item == "普通正文" and r.status == "success" for r in result.records))
        self.assertTrue(result.partial_success)
        self.assertAlmostEqual(formatted.sections[-1].top_margin.mm, 28, places=1)

    def test_word_refresh_corruption_keeps_readable_formatted_output(self):
        self.make_source()

        def damage_candidate(path):
            Path(path).write_bytes(b"not a readable document")
            raise RuntimeError("Word failed after modifying the file")

        self.refresh.side_effect = damage_candidate
        result, formatted = self.run_processor()
        self.assert_healthy_body_formatted(formatted)
        self.assert_preserved(formatted)
        self.assertTrue(any(r.item == "Word 目录及域自动刷新" and r.status == "skipped" for r in result.records))

    def test_cleanup_failure_keeps_readable_output_and_reports_skip(self):
        self.make_source()

        def damage_candidate(path, **kwargs):
            Path(path).write_bytes(b"broken cleanup candidate")
            raise RuntimeError("cleanup failed")

        with patch("word_formatter.core.processor.finalize_docx", side_effect=damage_candidate):
            result, formatted = self.run_processor()
        self.assert_healthy_body_formatted(formatted)
        self.assert_preserved(formatted)
        self.assertTrue(any(r.item == "最终安全清理" and r.status == "skipped" for r in result.records))
        self.assertFalse(any(r.item == "审阅批注" and r.status == "success" for r in result.records))

    def test_word_deleted_candidate_keeps_last_readable_output(self):
        self.make_source()

        def remove_candidate(path):
            Path(path).unlink()

        self.refresh.side_effect = remove_candidate
        result, formatted = self.run_processor()
        self.assert_healthy_body_formatted(formatted)
        self.assert_preserved(formatted)
        self.assertTrue(any(r.item == "Word 目录及域自动刷新" and r.status == "skipped" for r in result.records))

    def test_word_candidate_permission_error_does_not_revoke_output(self):
        self.make_source()
        self.refresh.side_effect = PermissionError("Word locked a private temporary path")
        result, formatted = self.run_processor()
        self.assert_healthy_body_formatted(formatted)
        self.assertTrue(any(r.item == "Word 目录及域自动刷新" and r.status == "skipped" for r in result.records))
        self.assertNotIn("private temporary path", " ".join(result.warnings))

    def test_cleanup_deleted_or_locked_candidate_preserves_saved_result(self):
        self.make_source()

        def remove_candidate(path, **kwargs):
            Path(path).unlink()
            return {}

        for fault in (remove_candidate, PermissionError("candidate is locked")):
            with self.subTest(fault=type(fault).__name__):
                with patch("word_formatter.core.processor.finalize_docx", side_effect=fault):
                    result, formatted = self.run_processor()
                self.assert_healthy_body_formatted(formatted)
                self.assertTrue(any(r.item == "最终安全清理" and r.status == "skipped" for r in result.records))

    def test_unreadable_input_remains_fatal(self):
        self.source.write_bytes(b"not a docx")
        with self.assertRaisesRegex(RuntimeError, "原文件未改动"):
            DocumentProcessor().process(self.source, self.rules, self.output)
        self.assertFalse(self.output.exists())

    def test_disk_write_failure_is_fatal_not_partial_success(self):
        self.make_source()
        original_save = DocumentObject.save

        def fail_disk(document, target):
            if not hasattr(target, "write"):
                raise OSError(28, "No space left on device")
            return original_save(document, target)

        with patch.object(DocumentObject, "save", fail_disk):
            with self.assertRaisesRegex(RuntimeError, "No space left"):
                DocumentProcessor().process(self.source, self.rules, self.output)
        self.assertFalse(self.output.exists())
        self.assertEqual(self.source.read_bytes(), self.original)

    def test_keyboard_interrupt_is_not_swallowed(self):
        self.make_source()
        with patch.object(DocumentProcessor, "_apply_headings", side_effect=KeyboardInterrupt):
            with self.assertRaises(KeyboardInterrupt):
                DocumentProcessor().process(self.source, self.rules, self.output)
        self.assertEqual(self.source.read_bytes(), self.original)

    def test_log_write_failure_does_not_discard_readable_docx(self):
        self.make_source()
        with patch.object(ProcessResult, "save_log", side_effect=PermissionError("private local log path")):
            result, formatted = self.run_processor()
        self.assert_healthy_body_formatted(formatted)
        self.assertTrue(any(r.item == "处理日志保存" and r.status == "skipped" for r in result.records))
        self.assertNotIn("private local log path", " ".join(result.warnings))

    def test_reference_exclusions_survive_module_checkpoint_restore(self):
        self.make_source()
        document = Document(self.source)
        document.add_heading("参考文献", level=1)
        document.add_paragraph("[1] 作者. 示例参考文献[J]. 测试, 2024.")
        document.save(self.source)
        self.original = self.source.read_bytes()
        self.rules.reference.enabled = True
        self.rules.figure_caption.enabled = True
        self.rules.reference.font_size_pt = 10
        with patch.object(DocumentProcessor, "_apply_figure_captions", side_effect=RuntimeError("optional caption failure")):
            result, formatted = self.run_processor()
        reference = next(p for p in formatted.paragraphs if p.text.startswith("[1]"))
        self.assertEqual(reference.runs[0].font.size.pt, 10)
        self.assert_healthy_body_formatted(formatted)
        self.assertTrue(result.partial_success)

    def test_process_result_legacy_success_count_is_compatible(self):
        result = ProcessResult(Path("source"), Path("output"))
        result.records.extend([
            ChangeRecord(None, "done", "", "", ""),
            ChangeRecord(None, "skip", "", "", "", status="skipped"),
        ])
        self.assertEqual(result.changed_count, 1)
        self.assertEqual(result.skipped_count, 1)
        self.assertTrue(result.partial_success)


if __name__ == "__main__":
    unittest.main()
