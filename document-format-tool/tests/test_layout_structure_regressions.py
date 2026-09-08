from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree

from word_formatter.core.analyzer import DocumentAnalyzer
from word_formatter.core.finalizer import finalize_docx
from word_formatter.core.processor import DocumentProcessor
from word_formatter.models.results import ProcessResult
from word_formatter.models.rules import DocumentRules


def result():
    return ProcessResult(Path("source.docx"), Path("output.docx"))


def complex_field(paragraph, instruction, cached):
    start = OxmlElement("w:fldChar")
    start.set(qn("w:fldCharType"), "begin")
    paragraph.add_run()._r.append(start)
    for part in (instruction[:3], instruction[3:]):
        node = OxmlElement("w:instrText")
        node.text = part
        paragraph.add_run()._r.append(node)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    paragraph.add_run()._r.append(separate)
    paragraph.add_run(cached)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph.add_run()._r.append(end)
    return start


class HeadingStructureTests(unittest.TestCase):
    def test_numbered_semantics_override_wrong_style_and_template_outline(self):
        document, rules = Document(), DocumentRules()
        rules.heading_2.enabled = rules.heading_3.enabled = True
        rules.heading_2.outline_level = 0
        rules.heading_3.outline_level = 0
        second = document.add_paragraph("1.1 研究意义", style="Heading 3")
        second._p.get_or_add_pPr().append(OxmlElement("w:rPr"))
        third = document.add_paragraph("1.1.1研究方法", style="Heading 1")
        self.assertEqual(DocumentAnalyzer.recognized_heading_level(second), 2)
        self.assertEqual(DocumentAnalyzer.recognized_heading_level(third), 3)
        DocumentProcessor._apply_headings(document, rules, result())
        self.assertEqual(second._p.pPr.find(qn("w:outlineLvl")).get(qn("w:val")), "1")
        self.assertEqual(third._p.pPr.find(qn("w:outlineLvl")).get(qn("w:val")), "2")
        self.assertEqual(second.style.name, "Heading 2")
        self.assertEqual(third.style.name, "Heading 3")
        children = [child.tag for child in second._p.pPr]
        self.assertLess(children.index(qn("w:outlineLvl")), children.index(qn("w:rPr")))
        self.assertLess(children.index(qn("w:spacing")), children.index(qn("w:outlineLvl")))

    def test_dates_cover_captions_and_appendix_programs_are_not_toc_members(self):
        document = Document()
        excluded = [document.add_paragraph(text, style="Heading 1") for text in ("2026年9月", "（2027）届毕业论文", "封面", "表1-1 参数配置", "图2-1 电路图")]
        for paragraph in excluded:
            self.assertIsNone(DocumentAnalyzer.recognized_heading_level(paragraph), paragraph.text)
        document.add_paragraph("附录一 程序清单", style="Heading 1")
        program = document.add_paragraph("（1）主控程序（ESP32）", style="Heading 3")
        references = document.add_paragraph("参考文献", style="Heading 1")
        DocumentProcessor._exclude_non_content_toc_entries(document, 1)
        for paragraph in excluded + [program]:
            self.assertEqual(paragraph._p.pPr.find(qn("w:outlineLvl")).get(qn("w:val")), "9")
        self.assertIsNone(references._p.pPr.find(qn("w:outlineLvl")))

    def test_normal_text_cannot_inherit_heading_outline_from_template(self):
        document, rules = Document(), DocumentRules()
        body = document.add_paragraph("这是一段必须保留且不能出现在目录中的作者正文。")
        rules.normal_text.outline_level = 0
        DocumentProcessor._apply_normal_text(document, rules.normal_text, result())
        self.assertEqual(body._p.pPr.find(qn("w:outlineLvl")).get(qn("w:val")), "9")

    def test_table_caption_default_is_centered_without_first_line(self):
        document, rules = Document(), DocumentRules()
        caption = document.add_paragraph("表1-1 测试参数", style="Heading 2")
        caption.paragraph_format.first_line_indent = Cm(1)
        DocumentProcessor._apply_table_captions(document, rules.table_caption, result())
        self.assertEqual(caption.alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(caption._p.pPr.find(qn("w:ind")).get(qn("w:firstLineChars")), "0")
        self.assertEqual(caption._p.pPr.find(qn("w:outlineLvl")).get(qn("w:val")), "9")


class TocStyleTests(unittest.TestCase):
    def test_styles_and_nested_hyperlink_runs_keep_confirmed_format(self):
        document, rules = Document(), DocumentRules()
        rules.toc_2.font_size_pt = 10.5
        rules.toc_2.chinese_font = "宋体"
        rules.toc_2.latin_font = "Times New Roman"
        rules.toc_2.left_indent_cm = 0.74
        rules.toc_2.special_indent_mode = "none"
        document.styles.add_style("TOC 2", WD_STYLE_TYPE.PARAGRAPH)
        sdt = OxmlElement("w:sdt")
        content = OxmlElement("w:sdtContent")
        paragraph = document.add_paragraph(style="TOC 2")
        link = OxmlElement("w:hyperlink")
        link.set(qn("w:anchor"), "_Toc1")
        run = OxmlElement("w:r")
        r_pr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:asciiTheme"), "minorHAnsi")
        r_pr.append(fonts)
        run.append(r_pr)
        text = OxmlElement("w:t")
        text.text = "1.1 Test\t2"
        run.append(text)
        link.append(run)
        paragraph._p.append(link)
        content.append(paragraph._p)
        sdt.append(content)
        document.element.body.insert(0, sdt)
        DocumentProcessor._apply_toc(document, rules, result())
        for r_pr in (run.find(qn("w:rPr")), document.styles["TOC 2"]._element.rPr):
            self.assertEqual(r_pr.find(qn("w:rFonts")).get(qn("w:ascii")), "Times New Roman")
            self.assertIsNone(r_pr.find(qn("w:rFonts")).get(qn("w:asciiTheme")))
            self.assertEqual(r_pr.find(qn("w:sz")).get(qn("w:val")), "21")
        self.assertEqual(document.styles["TOC 2"]._element.pPr.find(qn("w:ind")).get(qn("w:left")), str(Cm(0.74).twips))
        self.assertEqual(paragraph._p.pPr.find(qn("w:outlineLvl")).get(qn("w:val")), "9")

    def test_last_word_refresh_is_followed_by_toc_format_application(self):
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / "source.docx", Path(directory) / "out.docx"
            document = Document()
            document.add_paragraph("第一章 绪论", style="Heading 1")
            document.add_paragraph("作者正文必须保留。")
            document.save(source)
            rules = DocumentRules()
            rules.toc_1.font_size_pt = 10.5
            def refreshed(path):
                doc = Document(path)
                for p in doc.paragraphs:
                    if p.style.name.casefold() == "toc 1":
                        p.add_run("坏刷新格式").font.name = "Calibri"
                doc.save(path)
            with patch("word_formatter.core.processor.WordDocumentConverter.update_fields_in_place", side_effect=refreshed) as refresh:
                DocumentProcessor().process(source, rules, output)
            self.assertEqual(refresh.call_count, 2)
            doc = Document(output)
            toc = next(p for p in doc.paragraphs if p.style.name.casefold() == "toc 1")
            self.assertTrue(all(run.font.name == "Times New Roman" for run in toc.runs))
            self.assertEqual(doc.settings.element.find(qn("w:updateFields")).get(qn("w:val")), "false")


class ReferenceCacheTests(unittest.TestCase):
    def test_missing_refs_lock_cache_but_valid_cross_references_remain_live(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "refs.docx"
            document = Document()
            p = document.add_paragraph("保留引用[")
            complex_field(p, " REF missing \\h ", "6")
            p.add_run("]和[")
            complex_field(p, " REF valid \\h ", "9")
            p.add_run("]。")
            start = OxmlElement("w:bookmarkStart")
            start.set(qn("w:name"), "valid")
            start.set(qn("w:id"), "5")
            p._p.append(start)
            end = OxmlElement("w:bookmarkEnd")
            end.set(qn("w:id"), "5")
            p._p.append(end)
            simple = OxmlElement("w:fldSimple")
            simple.set(qn("w:instr"), ' REF "missing_simple" ')
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = "10"
            run.append(text)
            simple.append(run)
            p._p.append(simple)
            document.save(path)
            stats = finalize_docx(path)
            self.assertEqual(stats["unresolved_references_locked"], 2)
            with ZipFile(path) as package:
                root = etree.fromstring(package.read("word/document.xml"))
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            begins = root.xpath(".//w:fldChar[@w:fldCharType='begin']", namespaces=ns)
            self.assertEqual(begins[0].get(qn("w:fldLock")), "true")
            self.assertIsNone(begins[1].get(qn("w:fldLock")))
            self.assertEqual("".join(root.xpath(".//w:t/text()", namespaces=ns)), "保留引用[6]和[9]。10")
            self.assertEqual(finalize_docx(path)["unresolved_references_locked"], 0)


class SectionLayoutTests(unittest.TestCase):
    def test_toc_tail_merge_never_crosses_a_table_block(self):
        document = Document()
        document.add_paragraph("目录")
        document.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
        entry = document.add_paragraph("第一章 绪论\t1", "TOC 1")
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        document.add_paragraph().add_run()._r.append(end)
        document.add_table(1, 1).cell(0, 0).text = "不得越过的正文表格"
        document.add_section(WD_SECTION_START.NEW_PAGE)
        document.add_heading("第一章 绪论", 1)
        before = etree.tostring(document.element)
        DocumentProcessor._merge_toc_end_carriers(document)
        self.assertEqual(etree.tostring(document.element), before)

    def test_toc_end_and_section_tail_merge_without_losing_markers(self):
        document = Document()
        document.add_paragraph("目录")
        document.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
        entry = document.add_paragraph("第一章 绪论\t1", "TOC 1")
        tail = document.add_paragraph()
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        tail.add_run()._r.append(end)
        document.add_paragraph()
        document.add_section(WD_SECTION_START.NEW_PAGE)
        document.add_heading("第一章 绪论", 1)
        document.add_paragraph("必须保留的正文。")
        before_text = [p.text for p in document.paragraphs if p.text]
        DocumentProcessor._merge_toc_end_carriers(document)
        self.assertEqual([p.text for p in document.paragraphs if p.text], before_text)
        self.assertEqual(len(document.sections), 2)
        self.assertEqual(len(entry._p.xpath(".//w:fldChar[@w:fldCharType='end']")), 1)
        self.assertEqual(len(entry._p.xpath("./w:pPr/w:sectPr")), 1)
        self.assertEqual(entry._p.getnext().xpath(".//w:t/text()"), ["第一章 绪论"])
        DocumentProcessor._merge_toc_end_carriers(document)
        self.assertEqual(len(document.sections), 2)

    def test_invisible_toc_tail_and_break_carriers_compact_without_losing_structure(self):
        document = Document()
        document.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
        document.add_paragraph("参考文献\t35", style="TOC 1")
        field_end = document.add_paragraph()
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        field_end.add_run()._r.append(end)
        document.add_paragraph()
        document.add_section(WD_SECTION_START.NEW_PAGE)
        document.add_paragraph("第一章 绪论", style="Heading 1")
        document.add_paragraph("必须保留的作者文字。")
        empty_break = document.add_paragraph()
        empty_break.add_run().add_break()
        empty_break._p.xpath(".//w:br")[-1].set(qn("w:type"), "page")
        document.add_paragraph("结  论", style="Heading 1")
        before = (len(document.paragraphs), len(document.sections), len(document.element.body.xpath(".//w:fldChar")))
        DocumentProcessor._compact_layout_carriers(document)
        self.assertEqual(before, (len(document.paragraphs), len(document.sections), len(document.element.body.xpath(".//w:fldChar"))))
        for paragraph in [field_end, document.paragraphs[2], document.paragraphs[3], empty_break]:
            self.assertEqual(paragraph.paragraph_format.line_spacing.pt, 1)
            self.assertEqual(paragraph.paragraph_format.space_after.pt, 0)
            self.assertFalse(paragraph.paragraph_format.page_break_before)
        self.assertEqual(len(empty_break._p.xpath(".//w:br[@w:type='page']")), 1)

    def test_pdf_never_enters_word_conversion_even_with_invalid_cover_decision(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.docx"
            document = Document()
            document.add_paragraph("必须保留的原稿封面")
            document.add_paragraph("第一章 正文", style="Heading 1")
            document.save(source)
            with patch("word_formatter.core.processor.WordDocumentConverter.as_docx") as converter:
                revised = DocumentProcessor._compose_with_template_front(source, Path(directory) / "specification.pdf", result(), {"documentKind": "template", "copyFrontMatter": True})
            converter.assert_not_called()
            self.assertEqual(revised.paragraphs[0].text, "必须保留的原稿封面")

    def test_empty_break_before_conclusion_is_moved_without_removing_author_content(self):
        document = Document()
        document.add_paragraph("末章必须保留的正文。")
        empty_break = document.add_paragraph()
        empty_break.add_run().add_break()
        empty_break._p.xpath(".//w:br")[-1].set(qn("w:type"), "page")
        conclusion = document.add_paragraph("结  论", style="Heading 1")
        document.add_paragraph("完整的结论文字。")
        DocumentProcessor._start_chapters_on_new_pages(document, 1, result())
        self.assertEqual([p.text.strip() for p in document.paragraphs], ["末章必须保留的正文。", "结  论", "完整的结论文字。"])
        self.assertFalse(conclusion.paragraph_format.page_break_before)
        self.assertEqual(len(document.paragraphs[0]._p.xpath(".//w:br[@w:type='page']")), 1)

    def test_body_logical_pages_restart_without_inherited_page_footer(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.docx"
            output = Path(directory) / "output.docx"
            document = Document()
            document.add_paragraph("学校封面")
            document.add_paragraph("目录")
            document.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
            toc = document.add_paragraph(style="TOC 1")
            complex_field(toc, ' TOC \\o "1-3" \\h \\z \\u ', "第一章 绪论\t7")
            document.add_paragraph("第一章 绪论", style="Heading 1")
            document.add_paragraph("不得删除的作者正文。")
            document.save(source)
            rules = DocumentRules()
            self.assertFalse(rules.page_number.enabled)
            with patch("word_formatter.core.processor.WordDocumentConverter.update_fields_in_place"):
                DocumentProcessor().process(source, rules, output)
            revised = Document(output)
            self.assertEqual(len(revised.sections), 2)
            self.assertIsNone(revised.sections[0]._sectPr.find(qn("w:pgNumType")))
            self.assertEqual(revised.sections[1]._sectPr.find(qn("w:pgNumType")).get(qn("w:start")), "1")
            self.assertEqual(revised.sections[1]._sectPr.find(qn("w:pgNumType")).get(qn("w:fmt")), "decimal")
            self.assertFalse(revised.sections[0]._sectPr.findall(qn("w:footerReference")))
            self.assertFalse(revised.sections[1]._sectPr.findall(qn("w:footerReference")))
            self.assertIn("不得删除的作者正文。", [p.text for p in revised.paragraphs])

    def test_body_single_column_does_not_flatten_cover_sections(self):
        document = Document()
        document.add_paragraph("学校封面")
        document.add_section(WD_SECTION_START.NEW_PAGE)
        document.add_paragraph("目录")
        document.add_section(WD_SECTION_START.NEW_PAGE)
        body = document.add_paragraph("前言")
        document.add_paragraph("必须保留的作者正文。")
        for section in document.sections:
            section._sectPr.find(qn("w:cols")).set(qn("w:num"), "2")
        body_index = next(i for i, p in enumerate(document.paragraphs, 1) if p._p == body._p)
        self.assertEqual(DocumentProcessor._isolate_body_layout(document, body_index, result()), 0)
        self.assertEqual([s._sectPr.find(qn("w:cols")).get(qn("w:num")) for s in document.sections], ["2", "2", "1"])

    def test_new_content_boundary_keeps_prefix_columns_and_restarts_body_page_one(self):
        document = Document()
        document.add_paragraph("学校封面")
        document.add_paragraph("目录")
        document.add_paragraph("前言")
        document.add_paragraph("作者正文")
        document.sections[0]._sectPr.find(qn("w:cols")).set(qn("w:num"), "2")
        complex_field(document.sections[0].footer.paragraphs[0], " PAGE ", "1")
        added = DocumentProcessor._isolate_body_layout(document, 3, result())
        self.assertEqual(added, 1)
        DocumentProcessor._apply_page_numbers(document, {"normalize_existing": True}, result(), 3 + added)
        self.assertEqual(len(document.sections), 2)
        self.assertEqual(document.sections[0]._sectPr.find(qn("w:cols")).get(qn("w:num")), "2")
        self.assertEqual(document.sections[1]._sectPr.find(qn("w:cols")).get(qn("w:num")), "1")
        self.assertFalse(document.sections[0].footer._element.xpath(".//w:instrText"))
        self.assertEqual(document.sections[1]._sectPr.find(qn("w:pgNumType")).get(qn("w:start")), "1")
        self.assertTrue(document.sections[1].footer._element.xpath(".//w:instrText"))
        self.assertEqual([p.text for p in document.paragraphs if p.text], ["学校封面", "目录", "前言", "作者正文"])


if __name__ == "__main__":
    unittest.main()
