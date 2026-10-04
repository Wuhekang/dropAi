"""Derive the deployment layout asset from the accepted Word sample, never from a blank doc."""
import sys, hashlib, json
from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.oxml.ns import qn

source, target = map(Path, sys.argv[1:3])
doc = Document(source)
roles = {}
for p in doc.paragraphs:
    normalized = p.text.replace(' ', '').replace('\u3000', '')
    if normalized == '摘要': roles.setdefault('{{FRONT}}', deepcopy(p._p))
    if normalized.startswith('第1章') and p.style.name == 'Heading 1': roles.setdefault('{{CHAPTER}}', deepcopy(p._p))
    if p.text.startswith('1.1 ') and p.style.name == 'Heading 2': roles.setdefault('{{SECTION}}', deepcopy(p._p))
    if p.text.startswith('图1-1') or p.text.startswith('图 1-1'): roles.setdefault('{{CAPTION}}', deepcopy(p._p))
    if p.style.name == 'Title': roles.setdefault('{{TITLE}}', deepcopy(p._p))
    if p.style.name == 'Normal' and len(p.text) > 100: roles.setdefault('{{BODY}}', deepcopy(p._p))
if '{{TITLE}}' not in roles: roles['{{TITLE}}'] = deepcopy(doc.paragraphs[0]._p)
assert len(roles) == 6, list(roles)
baseline = [{'width': s.page_width, 'height': s.page_height, 'left': s.left_margin, 'right': s.right_margin} for s in doc.sections]
for element in list(doc._element.body):
    if element.tag != qn('w:sectPr'): doc._element.body.remove(element)
for role, element in roles.items():
    for child in list(element):
        if child.tag != qn('w:pPr'): element.remove(child)
    from docx.oxml import OxmlElement
    run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = role; run.append(text); element.append(run)
    doc._element.body.insert(len(doc._element.body)-1, element)
for rel_id, rel in list(doc.part.rels.items()):
    if rel.reltype.endswith('/image'): del doc.part.rels[rel_id]
doc.core_properties.author = 'Doki'
doc.core_properties.title = '机械论文排版母版'
target.parent.mkdir(parents=True, exist_ok=True)
doc.save(target)
print(json.dumps({'source_sha256': hashlib.sha256(source.read_bytes()).hexdigest(), 'roles': list(roles), 'sections': baseline, 'template_bytes': target.stat().st_size}, ensure_ascii=False))
